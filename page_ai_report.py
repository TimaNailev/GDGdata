import streamlit as st
import pandas as pd
import json
import os
import io
import re
from datetime import datetime
from docx import Document
from docx.shared import Inches
from fpdf import FPDF
import plotly.express as px
from ai_assistant import (
    get_report_template_suggestion,
    get_financial_insights,
)

# File-based storage paths
TEMPLATES_FILE = "report_templates.json"
HISTORY_FILE = "attribute_history.csv"

def load_templates():
    if not os.path.exists(TEMPLATES_FILE): return {}
    with open(TEMPLATES_FILE, "r") as f:
        try: return json.load(f)
        except: return {}

def save_templates(templates):
    with open(TEMPLATES_FILE, "w") as f:
        json.dump(templates, f, indent=4)

def load_history():
    if not os.path.exists(HISTORY_FILE):
        return pd.DataFrame(columns=["version_id", "attribute_name", "attribute_value", "timestamp"])
    return pd.read_csv(HISTORY_FILE)

def save_history(history_df):
    history_df.to_csv(HISTORY_FILE, index=False)

def safe_str(text):
    """Sanitize text for FPDF latin-1 encoded fonts."""
    if text is None: return ""
    text = str(text)
    text = text.replace('“', '"').replace('”', '"').replace('‘', "'").replace('’', "'").replace('–', '-')
    try: return text.encode('latin-1', 'replace').decode('latin-1')
    except: return "".join(c if ord(c) < 256 else '?' for c in text)

def clean_ai_text(text):
    """Clean common AI artifacts from Strategic Insights."""
    if not text: return ""
    text = re.sub(r"(?i)the\s+final\s+answer\s+is.*?\n", "", text)
    text = re.sub(r"(?i)##\s+Step\s+\d+:.*?\n", "", text)
    text = re.sub(r"(?i)\*\*Step\s+\d+:.*?\*\*\n", "", text)
    text = re.sub(r"(?i)^No\s+specific\s+numerical\s+answer.*?\n", "", text, flags=re.MULTILINE)
    # Extraction: Keep everything EXCEPT the JSON block for the textual insights
    match = re.search(r"(.*?)```json", text, re.DOTALL)
    if match:
        insight_part = match.group(1)
    else:
        insight_part = text
    
    insight_part = re.sub(r"\n{3,}", "\n\n", insight_part).strip()
    return safe_str(insight_part)

def truncate_text(text, limit=30):
    text = str(text)
    return (text[:limit] + '..') if len(text) > limit else text

def prepare_data_slice(df, sec):
    """Sort and limit dataframe based on section configuration."""
    sdf = df.copy()
    sort_by = sec.get('sort_by')
    sort_order = sec.get('sort_order', 'Descending')
    limit = sec.get('limit', 30)
    
    if sort_by and sort_by in sdf.columns:
        sdf = sdf.sort_values(by=sort_by, ascending=(sort_order == 'Ascending'))
    
    return sdf.head(limit)

# --- DOCX Generation ---
def generate_docx(df, conf, insights_text=None):
    doc = Document()
    if os.path.exists("logo.png"):
        table = doc.add_table(rows=1, cols=1)
        from docx.oxml.ns import nsdecls
        from docx.oxml import parse_xml
        shading_elm = parse_xml(r'<w:shd {} w:fill="143464"/>'.format(nsdecls('w')))
        table.rows[0].cells[0]._tc.get_or_add_tcPr().append(shading_elm)
        paragraph = table.rows[0].cells[0].paragraphs[0]
        paragraph.add_run().add_picture("logo.png", width=Inches(1.2))
    
    doc.add_heading(conf.get('report_name', 'Financial Report'), 0)
    clean_insights = clean_ai_text(insights_text)
    
    for sec in conf.get('sections', []):
        stype = sec.get('type')
        title = sec.get('title', '')
        if title: doc.add_heading(title, level=1)
            
        if stype == "summary_metrics":
            kpis, agg = sec.get('kpis', []), sec.get('agg', 'Sum')
            for k in kpis:
                if k in df.columns:
                    series = pd.to_numeric(df[k], errors='coerce')
                    val = series.sum() if agg == "Sum" else series.mean()
                    if pd.notnull(val):
                        doc.add_paragraph(f"{agg} of {k}: {val:,.2f}", style='List Bullet')
                    
        elif stype == "details_table":
            if sec.get('strategy'): doc.add_paragraph(f"Strategy: {sec['strategy']}", style='Caption')
            cols = sec.get('display_cols', [])
            valid_cols = [c for c in cols if c in df.columns][:10]
            if valid_cols:
                sdf = prepare_data_slice(df, sec)
                table = doc.add_table(rows=1, cols=len(valid_cols))
                table.style = 'Table Grid'
                for i, c in enumerate(valid_cols): table.rows[0].cells[i].text = c
                for i, row in sdf.iterrows():
                    cells = table.add_row().cells
                    for j, c in enumerate(valid_cols): cells[j].text = truncate_text(row[c], 20)
                        
        elif stype == "visual_chart":
            x_col, y_col = sec.get('x_col'), sec.get('y_col')
            if x_col in df.columns and y_col in df.columns:
                try:
                    # Sync chart with prioritized slice if applicable
                    sdf = prepare_data_slice(df, sec) if sec.get('sort_by') else df.head(20)
                    fig = px.bar(sdf, x=x_col, y=y_col, title=title)
                    img_path = f"tmp_docx_{datetime.now().microsecond}.png"
                    fig.write_image(img_path)
                    doc.add_picture(img_path, width=Inches(5.0))
                    if os.path.exists(img_path): os.remove(img_path)
                except: pass
        elif stype in ["ai_insights", "trend_analysis"]:
            if clean_insights: doc.add_paragraph(clean_insights)
            else: doc.add_paragraph("Analysis pending... Click 'Build Strategic Insights' in the generator.")
        elif stype == "text_block":
            doc.add_paragraph(sec.get('content', ''))

    doc.add_paragraph(f"\nReport Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# --- PDF Generation ---
def generate_pdf(df, conf, insights_text=None):
    pdf = FPDF()
    pdf.add_page()
    if os.path.exists("logo.png"):
        pdf.set_fill_color(20, 52, 100) 
        pdf.rect(155, 8, 45, 20, 'F')
        pdf.image("logo.png", x=160, y=10, w=35)
    
    pdf.set_font("Helvetica", 'B', 18)
    pdf.set_y(25)
    pdf.cell(190, 15, txt=safe_str(conf.get('report_name', 'Financial Report')), ln=True)
    
    clean_insights = clean_ai_text(insights_text)
    for sec in conf.get('sections', []):
        stype, title = sec.get('type'), safe_str(sec.get('title', ''))
        if title:
            pdf.ln(5)
            pdf.set_font("Helvetica", 'B', 12)
            pdf.cell(190, 10, txt=title.upper(), ln=True, border='B')
            pdf.ln(2)
            
        pdf.set_font("Helvetica", '', 10)
        if stype == "summary_metrics":
            kpis, agg = sec.get('kpis', []), sec.get('agg', 'Sum')
            for k in kpis:
                if k in df.columns:
                    series = pd.to_numeric(df[k], errors='coerce')
                    val = series.sum() if agg == "Sum" else series.mean()
                    if pd.notnull(val):
                        pdf.cell(190, 8, txt=safe_str(f"- {agg} of {k}: {val:,.2f}"), ln=True)
                    
        elif stype == "details_table":
            if sec.get('strategy'):
                pdf.set_font("Helvetica", 'I', 8)
                pdf.cell(190, 6, txt=safe_str(f"Strategy: {sec['strategy']}"), ln=True)
                pdf.set_font("Helvetica", '', 10)
            
            cols = sec.get('display_cols', [])
            valid_cols = [c for c in cols if c in df.columns][:8] 
            if valid_cols:
                sdf = prepare_data_slice(df, sec)
                pdf.set_font("Helvetica", 'B', 7)
                col_width = 190 / len(valid_cols)
                for c in valid_cols: pdf.cell(col_width, 8, txt=safe_str(truncate_text(c, 12)), border=1, align='C')
                pdf.ln()
                pdf.set_font("Helvetica", '', 6)
                for i, row in sdf.iterrows():
                    for c in valid_cols:
                        val = safe_str(truncate_text(row[c], 15))
                        pdf.cell(col_width, 6, txt=val, border=1)
                    pdf.ln()
        
        elif stype == "visual_chart":
            x_col, y_col = sec.get('x_col'), sec.get('y_col')
            if x_col in df.columns and y_col in df.columns:
                try:
                    sdf = prepare_data_slice(df, sec) if sec.get('sort_by') else df.head(20)
                    fig = px.bar(sdf, x=x_col, y=y_col)
                    img_path = f"tmp_pdf_{datetime.now().microsecond}.png"
                    fig.write_image(img_path)
                    pdf.image(img_path, w=140, x=25)
                    if os.path.exists(img_path): os.remove(img_path)
                except: pass
        elif stype in ["ai_insights", "trend_analysis"]:
            if clean_insights:
                pdf.set_font("Helvetica", '', 9)
                pdf.multi_cell(0, 6, txt=clean_insights.replace('**', ''))
            else:
                pdf.set_font("Helvetica", 'I', 9)
                pdf.cell(190, 8, txt="Analysis pending... Click 'Build Strategic Insights' in the generator.", ln=True)
        elif stype == "text_block":
            pdf.multi_cell(0, 7, txt=safe_str(sec.get('content', '')))
        
    pdf.ln(10)
    pdf.set_font("Helvetica", 'I', 8)
    pdf.cell(190, 10, txt=safe_str(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"), align='C')
    return bytes(pdf.output())

# --- Streamlit UI ---
def show_ai_financial_report(df):
    st.header("📈 AI Financial Management Report")
    if df is None:
        st.warning("⚠️ No data available.")
        return

    st.markdown("### Executive reporting Studio")
    tabs = st.tabs(["🎨 Template Studio", "📄 Report Generator", "📊 Historical Trends"])
    templates, history_df = load_templates(), load_history()
    
    with tabs[0]:
        st.subheader("Smart Template Builder")
        with st.expander("🎖️ AI Tactical Designer", expanded=True):
            user_goal = st.text_area("Reporting Objective", value="Professional Statement with prioritized Transaction History.")
            if st.button("Generate Tactical Draft", type="primary", use_container_width=True):
                with st.spinner("Drafting..."):
                    suggestion = get_report_template_suggestion(df, user_goal)
                    json_match = re.search(r"```json\s*\n(.*?)\n```", suggestion, re.DOTALL)
                    if json_match:
                        try: st.session_state["pending_template"] = json.loads(json_match.group(1))
                        except: st.error("JSON Error.")
                    else: st.write(suggestion)

        if "pending_template" in st.session_state:
            st.json(st.session_state["pending_template"])
            ccname = st.text_input("Template Name", value=st.session_state["pending_template"].get("report_name", "Report"))
            if st.button("Save & Confirm Template", use_container_width=True, type="primary"):
                templates[ccname] = st.session_state["pending_template"]
                save_templates(templates)
                del st.session_state["pending_template"]
                st.rerun()

        if templates:
            st.subheader("Stored Templates")
            for name, conf in templates.items():
                with st.expander(f"📋 {name}"):
                    st.write(conf)
                    if st.button("Delete", key=f"del_{name}"):
                        del templates[name]
                        save_templates(templates)
                        st.rerun()

    with tabs[1]:
        st.subheader("Executive Report Review")
        if not templates: st.warning("No templates found.")
        else:
            sel_name = st.selectbox("Select Template", list(templates.keys()))
            conf = templates[sel_name]
            
            if st.button("🧠 Build Strategic Insights & Recommendations", type="primary", use_container_width=True):
                with st.spinner("Analyzing..."):
                    raw_insights = get_financial_insights(df, json.dumps(conf))
                    json_match = re.search(r"```json\s*\n(.*?)\n```", raw_insights, re.DOTALL)
                    if json_match:
                        try:
                            suggested = json.loads(json_match.group(1))
                            if suggested:
                                for c in suggested:
                                    # Inject new sections if they don't exist
                                    if not any(s.get('title') == c.get('title') for s in conf.setdefault('sections', [])):
                                        conf['sections'].append({"type": "visual_chart", **c})
                                templates[sel_name] = conf
                                save_templates(templates)
                        except: pass
                    st.session_state["current_insights"] = clean_ai_text(raw_insights)
                    st.rerun() # Force UI to see new charts and insights
            
            if "current_insights" in st.session_state:
                with st.chat_message("assistant", avatar="🎖️"): 
                    st.markdown(st.session_state["current_insights"])
            elif any(s.get('type') == 'ai_insights' for s in conf.get('sections', [])):
                st.info("💡 Click the button above to generate a strategic analysis for this template.")

            st.divider()
            with st.container(border=True):
                c1, c2 = st.columns([5, 1])
                with c1: st.title(f"{conf.get('report_name', 'Report')}")
                with c2: 
                    if os.path.exists("logo.png"): st.image("logo.png", width=120)

                for sec in conf.get('sections', []):
                    stype, title = sec.get('type'), sec.get('title', '')
                    if title: st.subheader(title)
                    if stype == "summary_metrics":
                        kpis, agg = sec.get('kpis', []), sec.get('agg', 'Sum')
                        mcols = st.columns(len(kpis)) if kpis else [st]
                        for i, k in enumerate(kpis):
                            if k in df.columns:
                                series = pd.to_numeric(df[k], errors='coerce')
                                val = series.sum() if agg == "Sum" else series.mean()
                                if pd.notnull(val):
                                    mcols[i].metric(label=f"{agg} of {k}", value=f"{val:,.2f}")
                                else:
                                    mcols[i].error(f"{k} isn't numeric")
                    elif stype == "details_table":
                        if sec.get('strategy'): st.caption(f"Strategy: {sec['strategy']}")
                        sdf = prepare_data_slice(df, sec)
                        st.dataframe(sdf, use_container_width=True)
                    elif stype == "visual_chart":
                        x, y = sec.get('x_col'), sec.get('y_col')
                        if x in df.columns and y in df.columns:
                            try:
                                sdf = prepare_data_slice(df, sec) if sec.get('sort_by') else df.head(20)
                                st.plotly_chart(px.bar(sdf, x=x, y=y), use_container_width=True)
                            except: st.error("Chart Fail")
                    elif stype in ["ai_insights", "trend_analysis"]:
                        if "current_insights" in st.session_state: st.markdown(st.session_state["current_insights"])

            st.divider()
            c1, c2 = st.columns(2)
            with c1:
                doc_buf = generate_docx(df, conf, st.session_state.get("current_insights"))
                st.download_button("📥 Word (.docx)", doc_buf, f"{sel_name}.docx", use_container_width=True)
            with c2:
                try:
                    pdf_data = generate_pdf(df, conf, st.session_state.get("current_insights"))
                    st.download_button("📥 PDF (.pdf)", pdf_data, f"{sel_name}.pdf", type="primary", use_container_width=True)
                except Exception as e: st.error(f"PDF Error: {e}")

    with tabs[2]:
        st.subheader("Historical trends")
        if history_df.empty: st.info("No snapshots found.")
        else:
            all_attrs = history_df['attribute_name'].unique().tolist()
            sel_plot = st.multiselect("Trend selection", all_attrs, default=all_attrs[:1])
            if sel_plot:
                pdf = history_df[history_df['attribute_name'].isin(sel_plot)]
                st.plotly_chart(px.line(pdf, x="version_id", y="attribute_value", color="attribute_name", markers=True), use_container_width=True)
